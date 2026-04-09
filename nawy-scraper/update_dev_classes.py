#!/usr/bin/env python3
"""Update developers with their class (A/B/C) from nawydev_reranked.xlsx."""
import sys
if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

import openpyxl
from database.connection import db_client

db_client.connect()
col = db_client._db['developers']

# Read Excel
wb = openpyxl.load_workbook('nawydev_reranked.xlsx')
ws = wb.active

# Build lookup: developer name -> class
class_lookup = {}
for row in range(2, ws.max_row + 1):
    name = ws.cell(row, 1).value
    dev_class = ws.cell(row, 2).value
    if name and dev_class:
        class_lookup[name.strip().lower()] = dev_class.strip()

print(f"Loaded {len(class_lookup)} developer classes from Excel")

# Update MongoDB
updated = 0
not_found = []
for dev in col.find({}, {'_id': 1, 'dev_name': 1}):
    dev_name = (dev.get('dev_name') or '').strip().lower()
    if dev_name in class_lookup:
        col.update_one(
            {'_id': dev['_id']},
            {'$set': {'developer_class': class_lookup[dev_name]}}
        )
        updated += 1
    else:
        not_found.append(dev.get('dev_name', ''))

print(f"Updated {updated} developers with class")
print(f"Not matched: {len(not_found)}")

db_client.disconnect()
